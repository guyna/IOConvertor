#!/usr/bin/env python3
"""Convert TITAN IOC exports into CrowdStrike, SentinelOne and Rapid7 files."""

from __future__ import annotations

import csv
import ipaddress
import re
from collections import OrderedDict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple

CS_FIELDS = [
    "type",
    "value",
    "action",
    "severity",
    "description",
    "platforms",
    "applied_globally",
    "expiration",
]
S1_FIELDS = [
    "OS",
    "Description",
    "SHA1",
    "SHA256",
    "Scope",
    "Scope Path",
    "User",
    "Last Update",
    "Source",
]
S1_OS = ("windows", "linux", "osx")
S1_SCOPE = "site"
S1_SCOPE_PATH = "Global"  # default; overridable from the UI
S1_SOURCE = "user"
CS_PLATFORMS = "windows,mac,linux"

SHA256_RE = re.compile(r"^[a-fA-F0-9]{64}$")
SHA1_RE = re.compile(r"^[a-fA-F0-9]{40}$")
MD5_RE = re.compile(r"^[a-fA-F0-9]{32}$")
IPV4_RE = re.compile(
    r"\b(?:(?:25[0-5]|2[0-4]\d|[01]?\d\d?)\.){3}(?:25[0-5]|2[0-4]\d|[01]?\d\d?)\b"
)

SUPPORTED_SUFFIXES = {".csv", ".xlsx", ".xlsb"}


def add_months(d: date, months: int = 3) -> date:
    month = d.month - 1 + months
    year = d.year + month // 12
    month = month % 12 + 1
    day = min(d.day, _days_in_month(year, month))
    return date(year, month, day)


def _days_in_month(year: int, month: int) -> int:
    import calendar

    return calendar.monthrange(year, month)[1]


def expiration_iso(today: Optional[date] = None) -> str:
    d = add_months(today or date.today(), 3)
    return f"{d.isoformat()}T00:00:00Z"


def undefang(value: str) -> str:
    if not value:
        return ""
    s = value.strip()
    s = s.replace("hxxps://", "https://").replace("HXXPS://", "https://")
    s = s.replace("hxxp://", "http://").replace("HXXP://", "http://")
    s = s.replace("hxxps:", "https:").replace("hxxp:", "http:")
    s = s.replace("[.]", ".").replace("[dot]", ".").replace("(.)", ".")
    s = s.replace("[:]", ":").replace("[://]", "://")
    s = s.replace("[@]", "@")
    s = s.replace("[", "").replace("]", "")
    return s.strip()


def ensure_url_protocol(url: str) -> str:
    u = undefang(url)
    if not u:
        return ""
    low = u.lower()
    if low.startswith(("http://", "https://", "ftp://")):
        return u
    return "https://" + u


def is_ip(value: str) -> bool:
    try:
        ipaddress.ip_address(value)
        return True
    except ValueError:
        return False


def extract_ip_from_url(url: str) -> Optional[str]:
    u = undefang(url)
    host = ""
    m = re.match(r"^[a-zA-Z][a-zA-Z0-9+.-]*://([^/]+)", u)
    if m:
        host = m.group(1)
    else:
        return None
    host = host.split("@")[-1]
    if host.startswith("["):
        end = host.find("]")
        if end > 0:
            host = host[1:end]
    else:
        host = host.split(":")[0]
    host = host.strip().strip(".")
    return host if host and is_ip(host) else None


def norm_header(name: str) -> str:
    return (name or "").strip().lower().replace(" ", "_")


def _cell_str(val: Any) -> str:
    if val is None:
        return ""
    if isinstance(val, bytes):
        val = val.decode("utf-8", "replace")
    if isinstance(val, float) and val.is_integer():
        return str(int(val))
    if isinstance(val, int):
        return str(val)
    text = str(val).strip()
    if text.endswith(".0") and text.replace(".", "", 1).isdigit():
        return text[:-2]
    return text


def row_get(row: Dict[str, Any], *keys: str) -> str:
    lowered = {norm_header(k): v for k, v in row.items()}
    for key in keys:
        text = _cell_str(lowered.get(norm_header(key), ""))
        if text:
            return text
    return ""


def normalize_type(raw: str) -> str:
    t = (raw or "").strip().lower()
    aliases = {
        "sha-256": "sha256",
        "sha_256": "sha256",
        "sha256sum": "sha256",
        "sha-1": "sha1",
        "sha_1": "sha1",
        "ip": "ipv4",
        "ip-dst": "ipv4",
        "ip-src": "ipv4",
        "ip_dst": "ipv4",
        "ip_src": "ipv4",
        "ipv4-addr": "ipv4",
        "ipv6": "ipv6",
        "ipv6-addr": "ipv6",
        "ip-dst|port": "ipv4",
        "hostname": "domain",
        "domain-name": "domain",
        "fqdn": "domain",
        "md5sum": "md5",
        "uri": "url",
        "link": "url",
    }
    return aliases.get(t, t)


def alert_label(event_id: str, event_info: str) -> str:
    eid = (event_id or "").strip()
    info = (event_info or "").strip()
    if eid and info:
        return f"Alert {eid} - {info}"
    if eid:
        return f"Alert {eid}"
    if info:
        return info
    return "Alert TITAN"


def read_csv_rows(path: Path) -> List[Dict[str, Any]]:
    for enc in ("utf-8-sig", "utf-8", "cp1255", "latin-1"):
        try:
            with path.open("r", encoding=enc, newline="") as fh:
                sample = fh.read(4096)
                fh.seek(0)
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
                except csv.Error:
                    dialect = csv.excel
                reader = csv.DictReader(fh, dialect=dialect)
                return [dict(r) for r in reader]
        except UnicodeDecodeError:
            continue
    return []


def read_xlsx_rows(path: Path) -> List[Dict[str, Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    headers = [str(h) if h is not None else "" for h in next(rows_iter, [])]
    out: List[Dict[str, Any]] = []
    for raw in rows_iter:
        if raw is None or all(c is None or str(c).strip() == "" for c in raw):
            continue
        out.append({headers[i]: raw[i] if i < len(raw) else "" for i in range(len(headers))})
    wb.close()
    return out


def read_xlsb_rows(path: Path) -> List[Dict[str, Any]]:
    from pyxlsb import open_workbook

    with open_workbook(str(path)) as wb:
        sheet = wb.get_sheet(1)
        rows = list(sheet.rows())
    if not rows:
        return []
    headers = [str(c.v) if c.v is not None else "" for c in rows[0]]
    out: List[Dict[str, Any]] = []
    for raw in rows[1:]:
        vals = [c.v for c in raw]
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        out.append({headers[i]: vals[i] if i < len(vals) else "" for i in range(len(headers))})
    return out


def load_titan_file(path: Path) -> List[Dict[str, Any]]:
    suf = path.suffix.lower()
    if suf == ".csv":
        return read_csv_rows(path)
    if suf == ".xlsx":
        return read_xlsx_rows(path)
    if suf == ".xlsb":
        return read_xlsb_rows(path)
    raise ValueError(f"Unsupported file type: {path.suffix}")


def iter_indicators(rows: Iterable[Dict[str, Any]]) -> Iterable[Tuple[str, str, str]]:
    """Yield (norm_type, clean_value, description)."""
    for row in rows:
        raw_type = row_get(row, "type")
        raw_val = row_get(row, "value")
        event_id = row_get(row, "event_id")
        event_info = row_get(row, "event_info")
        desc = alert_label(event_id, event_info)
        ntype = normalize_type(raw_type)
        clean = undefang(str(raw_val))
        if ntype == "ip-dst|port" or "|" in clean and ntype in {"ipv4", "ipv6"}:
            clean = clean.split("|", 1)[0].strip()
        if not clean:
            continue
        yield ntype, clean, desc


def merge_desc(existing: str, incoming: str) -> str:
    parts = [p.strip() for p in existing.split(";") if p.strip()]
    if incoming and incoming not in parts:
        parts.append(incoming)
    # keep alert numbers stable
    def sort_key(s: str) -> Tuple[int, str]:
        m = re.search(r"Alert\s+(\d+)", s)
        return (int(m.group(1)) if m else 10**9, s)

    parts = sorted(dict.fromkeys(parts), key=sort_key)
    return "; ".join(parts)


class ConvertResult:
    def __init__(self) -> None:
        self.cs: "OrderedDict[Tuple[str, str], Dict[str, str]]" = OrderedDict()
        self.s1_sha256: "OrderedDict[str, str]" = OrderedDict()
        self.s1_sha1: "OrderedDict[str, str]" = OrderedDict()
        self.r7: "OrderedDict[str, None]" = OrderedDict()
        self.files_ok = 0
        self.files_fail: List[str] = []
        self.skipped = 0
        self.warnings: List[str] = []

    def add_cs(self, typ: str, value: str, action: str, desc: str, exp: str) -> None:
        key = (typ, value.lower())
        if key in self.cs:
            self.cs[key]["description"] = merge_desc(self.cs[key]["description"], desc)
            return
        self.cs[key] = {
            "type": typ,
            "value": value,
            "action": action,
            "severity": "HIGH",
            "description": desc,
            "platforms": CS_PLATFORMS,
            "applied_globally": "TRUE",
            "expiration": exp,
        }

    def add_s1_sha256(self, value: str, desc: str) -> None:
        k = value.lower()
        if k in self.s1_sha256:
            self.s1_sha256[k] = merge_desc(self.s1_sha256[k], desc)
        else:
            self.s1_sha256[k] = desc

    def add_s1_sha1(self, value: str, desc: str) -> None:
        k = value.lower()
        if k in self.s1_sha1:
            self.s1_sha1[k] = merge_desc(self.s1_sha1[k], desc)
        else:
            self.s1_sha1[k] = desc

    def add_r7(self, value: str) -> None:
        if value:
            self.r7.setdefault(value, None)


def convert_files(paths: List[Path], today: Optional[date] = None) -> ConvertResult:
    result = ConvertResult()
    exp = expiration_iso(today)
    for path in paths:
        try:
            rows = load_titan_file(path)
            result.files_ok += 1
        except Exception as exc:  # noqa: BLE001
            result.files_fail.append(f"{path.name}: {exc}")
            continue
        for ntype, value, desc in iter_indicators(rows):
            if ntype == "sha256":
                if not SHA256_RE.match(value):
                    result.skipped += 1
                    continue
                hv = value.lower()
                result.add_cs("sha256", hv, "prevent", desc, exp)
                result.add_s1_sha256(hv, desc)
            elif ntype == "sha1":
                if not SHA1_RE.match(value):
                    result.skipped += 1
                    continue
                result.add_s1_sha1(value.lower(), desc)
            elif ntype == "md5":
                if not MD5_RE.match(value):
                    result.skipped += 1
                    continue
                result.add_r7(value.lower())
            elif ntype == "domain":
                host = value.lower().rstrip(".")
                if not host or is_ip(host):
                    if host and is_ip(host):
                        ip_type = "ipv6" if ":" in host else "ipv4"
                        result.add_cs(ip_type, host, "detect", desc, exp)
                        result.add_r7(host)
                    else:
                        result.skipped += 1
                    continue
                result.add_cs("domain", host, "detect", desc, exp)
                result.add_r7(host)
            elif ntype in {"ipv4", "ipv6"}:
                ip = value.split("/")[0].strip()
                if not is_ip(ip):
                    result.skipped += 1
                    continue
                ip_type = "ipv6" if ":" in ip else "ipv4"
                result.add_cs(ip_type, ip, "detect", desc, exp)
                result.add_r7(ip)
            elif ntype == "url":
                url = ensure_url_protocol(value)
                if url:
                    result.add_r7(url)
                ip = extract_ip_from_url(value)
                if ip:
                    ip_type = "ipv6" if ":" in ip else "ipv4"
                    result.add_cs(ip_type, ip, "detect", desc, exp)
            else:
                result.skipped += 1
    return result


ALL_TARGETS = ("cs", "s1", "r7")


def write_outputs(
    result: ConvertResult,
    out_dir: Path,
    stamp: Optional[str] = None,
    targets: Optional[Iterable[str]] = None,
    s1_scope_path: Optional[str] = None,
) -> Dict[str, Path]:
    """Write only the requested vendor files (default: all three)."""
    out_dir.mkdir(parents=True, exist_ok=True)
    stamp = stamp or date.today().isoformat()
    wanted = {t.lower() for t in (targets or ALL_TARGETS)} & set(ALL_TARGETS)
    if not wanted:
        raise ValueError("No output format selected.")
    scope_path = (s1_scope_path or "").strip() or S1_SCOPE_PATH
    all_files = {
        "cs": out_dir / f"CS_IOCs_{stamp}.csv",
        "s1": out_dir / f"S1_IOCs_{stamp}.csv",
        "r7": out_dir / f"R7_IOCs_{stamp}.csv",
    }
    files = {k: v for k, v in all_files.items() if k in wanted}
    if "cs" in files:
        _write_cs(result, files["cs"])
    if "s1" in files:
        _write_s1(result, files["s1"], scope_path)
    if "r7" in files:
        _write_r7(result, files["r7"])
    return files


def _write_cs(result: ConvertResult, path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=CS_FIELDS)
        w.writeheader()
        w.writerows(result.cs.values())


def _write_s1(result: ConvertResult, path: Path, scope_path: str) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=S1_FIELDS)
        w.writeheader()
        for sha, desc in result.s1_sha256.items():
            for os_name in S1_OS:
                w.writerow(
                    {
                        "OS": os_name,
                        "Description": desc,
                        "SHA1": "",
                        "SHA256": sha,
                        "Scope": S1_SCOPE,
                        "Scope Path": scope_path,
                        "User": "",
                        "Last Update": "",
                        "Source": S1_SOURCE,
                    }
                )
        for sha, desc in result.s1_sha1.items():
            for os_name in S1_OS:
                w.writerow(
                    {
                        "OS": os_name,
                        "Description": desc,
                        "SHA1": sha,
                        "SHA256": "",
                        "Scope": S1_SCOPE,
                        "Scope Path": scope_path,
                        "User": "",
                        "Last Update": "",
                        "Source": S1_SOURCE,
                    }
                )


def _write_r7(result: ConvertResult, path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.writer(fh)
        for val in result.r7.keys():
            w.writerow([val])
